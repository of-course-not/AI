# -*- coding: utf-8 -*-
"""实验二自检：.pkt 配置项 + 报告 docx + 教程 md + 配图 + IP 学号一致性。

内网必须是 192.168.90.0/24 的子网，公网必须是 200.0.90.0/24 的子网，
即每个 IPv4 地址的第 3 个字节必须是 90（学号后两位）。
"""
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
EXP2 = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import docx
import pktlib
from PIL import Image

PKT = os.path.join(EXP2, "03-工程文件")
REPORT = os.path.join(EXP2, "01-实验报告", "90-实验二-ACL与NAT配置-实验报告.docx")
TUTORIAL = os.path.join(EXP2, "02-新手教程",
                        "网络设备配置与管理-实验二（ACL与NAT）Packet Tracer 新手教程.md")
FIGS = os.path.join(EXP2, "06-图表与截图")

XX = "90"
out = []


def P(s=""):
    out.append(s)


def devices(xml):
    return re.findall(r"<DEVICE>(.*?)</DEVICE>", xml, re.S)


def nm(d):
    m = re.search(r"<NAME[^>]*>([^<]*)</NAME>", d)
    return m.group(1) if m else "?"


def run_of(d):
    m = re.search(r"<RUNNINGCONFIG>(.*?)</RUNNINGCONFIG>", d, re.S)
    if not m:
        return ""
    return "\n".join(re.findall(r"<LINE>(.*?)</LINE>", m.group(1), re.S))


def port_ips(d):
    res = []
    for p in re.findall(r"<PORT>(.*?)</PORT>", d, re.S):
        ip = re.search(r"<IP>([^<]*)</IP>", p)
        res.append(ip.group(1) if ip else "")
    return res


def check_pkt(fn, expect):
    path = os.path.join(PKT, fn)
    xml = pktlib.load(path).decode("utf-8", errors="replace")
    ds = devices(xml)
    nl = len(re.findall(r"<LINK>", xml))
    names = {nm(d): d for d in ds}
    ok = True
    for dev, needles in expect.items():
        d = names.get(dev)
        if d is None:
            P("   [缺设备] %s" % dev)
            ok = False
            continue
        run = run_of(d)
        for n in needles:
            if n not in run:
                P("   [缺配置] %-12s %s" % (dev, n))
                ok = False
    ver = re.search(r"<VERSION>([^<]*)</VERSION>", xml)
    P("%-38s %7d B  dev=%-3d link=%-3d ver=%-12s %s"
      % (fn, os.path.getsize(path), len(ds), nl,
         ver.group(1) if ver else "?", "OK" if ok else "FAIL"))
    return ok


SKIP_LINE = re.compile(r"router-id|area \d|encapsulation|clock rate|mac-address")


def bad_ips(text, where):
    bad = []
    for line in text.split("\n"):
        if SKIP_LINE.search(line):
            continue
        for ip in re.findall(r"\b(\d{1,3})\.(\d{1,3})\.(\d{1,3})\.(\d{1,3})\b", line):
            a, b, c, dd = ip
            if a in ("255", "0", "127", "169"):
                continue
            if not (a == "192" and b == "168") and not (a == "200" and b == "0"):
                continue
            if c != XX:
                bad.append("%s.%s.%s.%s  (%s)" % (a, b, c, dd, where))
    return bad


def strip_swap_section(text):
    """去掉教程里「换学号怎么改」这一小节——那里故意用别的数字做示例。"""
    i = text.find("换学号怎么改")
    if i < 0:
        return text
    j = text.find("\n---", i)
    return text[:i] + (text[j:] if j > 0 else "")


def check_ips():
    bad = []
    for fn in sorted(os.listdir(PKT)):
        if not fn.endswith(".pkt") or fn.startswith("_"):
            continue
        xml = pktlib.load(os.path.join(PKT, fn)).decode("utf-8", errors="replace")
        for d in devices(xml):
            name = nm(d)
            for ip in port_ips(d):
                if ip:
                    bad += bad_ips(ip, "%s %s 端口" % (fn, name))
            bad += bad_ips(run_of(d), "%s %s 配置" % (fn, name))

    d = docx.Document(REPORT)
    txt = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    bad += bad_ips(txt, "报告 docx")

    md = strip_swap_section(open(TUTORIAL, encoding="utf-8").read())
    bad += bad_ips(md, "教程 md")

    if bad:
        P("   [IP 不一致] 共 %d 处：" % len(bad))
        for x in bad[:40]:
            P("      " + x)
        return False
    P("   所有 IPv4 地址第 3 字节均为 %s（192.168.%s.x / 200.0.%s.x）  OK"
      % (XX, XX, XX))
    return True


P("=" * 84)
P("A. 实验二 2 个 .pkt 的配置项")
P("=" * 84)
a = True
a &= check_pkt("90-实验二一-ACL配置.pkt", {
    "CoreSW": ["ip routing", "interface Vlan2", "interface Vlan3", "interface Vlan4",
               "ip access-group 120 in", "ip access-group 130 in",
               "ip access-group 110 out", "access-list 110", "access-list 120",
               "access-list 130", "no switchport",
               "ip route 0.0.0.0 0.0.0.0 192.168.90.194"],
    "Router0": ["ip route 192.168.90.0 255.255.255.0 192.168.90.193",
                "ip route 0.0.0.0 0.0.0.0 200.0.90.2"],
    "Router1": ["interface Loopback0", "200.0.90.100"],
    "student": ["switchport access vlan 3", "switchport mode trunk"],
    "Server": ["switchport access vlan 4", "switchport mode trunk"],
})
a &= check_pkt("90-实验二二-NAT配置.pkt", {
    "NAT路由器": ["ip nat inside", "ip nat outside", "ip nat pool PUB",
               "ip nat inside source list 1 pool PUB",
               "ip nat inside source static tcp", "access-list 1 permit"],
    "Internet路由器": ["ip route 200.0.90.16 255.255.255.240 200.0.90.1"],
    "CoreSW": ["ip routing", "ip route 0.0.0.0 0.0.0.0 192.168.90.194"],
    "外网网络": ["vlan 10", "switchport access vlan 10"],
})

P()
P("=" * 84)
P("B. 学号后两位 XX=%s 的 IP 一致性" % XX)
P("=" * 84)
b = check_ips()

P()
P("=" * 84)
P("C. 报告 docx / 教程 md / 配图")
P("=" * 84)
d = docx.Document(REPORT)
s = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
must = ["ACL", "NAT", "192.168.90", "200.0.90", "ip access-group", "ip nat",
        "ip nat inside source static tcp", "no switchport", "Vlan4"]
miss = [m for m in must if m not in s]
P("报告 %8d B  表格=%d 段落=%d 图=%d 正文=%d 字  %s"
  % (os.path.getsize(REPORT), len(d.tables), len(d.paragraphs),
     len(d.inline_shapes), len(s), "OK" if not miss else "缺: " + ",".join(miss)))
c1 = not miss

t = open(TUTORIAL, encoding="utf-8").read()
fences = t.count("```")
P("教程 %8d B  %5d 行  H2=%-3d 围栏=%-4d %s"
  % (os.path.getsize(TUTORIAL), t.count("\n"), t.count("\n## "), fences,
     "OK" if fences % 2 == 0 else "围栏不成对"))
c2 = fences % 2 == 0

for fn in ("fig_overview2.png", "fig1_acl.png", "fig2_nat.png"):
    p = os.path.join(FIGS, fn)
    im = Image.open(p)
    P("   %-20s %-10s %8d B" % (fn, "%dx%d" % im.size, os.path.getsize(p)))

P()
P("=" * 84)
P("D. 汇总")
P("=" * 84)
P("配置项   : " + ("ALL OK" if a else "HAS FAILURES"))
P("IP=90    : " + ("ALL OK" if b else "HAS FAILURES"))
P("报告/教程: " + ("ALL OK" if (c1 and c2) else "HAS FAILURES"))

print("\n".join(out))
