# -*- coding: utf-8 -*-
"""Shared helpers for building the experiment report DOCX."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
CODE_FONT = "Consolas"


def _set_run_font(run, font=BODY_FONT, size=12, bold=False, color=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), font)


def shade(element, fill):
    ppr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_para(container, text="", size=12, bold=False, font=BODY_FONT, align=None,
             indent=True, space_before=0, space_after=3, line=1.4, color=None):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if align == "center":
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Pt(0)
    elif align == "right":
        pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "left":
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if text:
        run = p.add_run(text)
        _set_run_font(run, font=font, size=size, bold=bold, color=color)
    return p


def add_heading(container, text, level=1):
    sizes = {1: 16, 2: 14, 3: 12}
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(10 if level == 1 else 8)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.3
    pf.keep_with_next = True
    run = p.add_run(text)
    _set_run_font(run, font=HEAD_FONT, size=sizes.get(level, 12), bold=True)
    return p


def add_code(container, text):
    lines = text.strip("\n").split("\n")
    for idx, line in enumerate(lines):
        p = container.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.05
        pf.left_indent = Cm(0.5)
        pf.keep_with_next = idx < len(lines) - 1
        shade(p._p, "F4F4F4")
        run = p.add_run(line if line else " ")
        _set_run_font(run, font=CODE_FONT, size=9)
    # trailing spacer
    spacer = container.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 1.0
    return spacer


def add_caption(container, text):
    p = container.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    _set_run_font(run, font=BODY_FONT, size=9)
    return p


def _set_cell_margins(table, top=40, bottom=40, left=80, right=80):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tbl_pr.append(borders)


def add_table(container, header, rows, widths=None, size=9, header_fill="DCE6F1"):
    table = container.add_table(rows=1, cols=len(header))
    try:
        table.style = "Table Grid"
    except KeyError:
        _set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_cell_margins(table)

    hdr = table.rows[0]
    for i, text in enumerate(header):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(text)
        _set_run_font(run, font=HEAD_FONT, size=size, bold=True)
        shade(cell._tc.get_or_add_tcPr(), header_fill)

    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(text))
            _set_run_font(run, font=BODY_FONT, size=size)

    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)

    # keep the last row together with the caption printed below the table
    for cell in table.rows[-1].cells:
        for p in cell.paragraphs:
            p.paragraph_format.keep_with_next = True
    return table


def add_picture(container, path, width_cm=15.0):
    p = container.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Pt(0)
    p.add_run().add_picture(path, width=Cm(width_cm))
    return p


def setup_document(path):
    doc = Document(path)
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), BODY_FONT)
    rfonts.set(qn("w:hAnsi"), BODY_FONT)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    return doc


def set_cell_lines(cell, lines, size=10.5, font=BODY_FONT, bold_first=False):
    """Fill a cell: first paragraph exists already, extra lines appended."""
    cell.text = ""
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.35
        run = p.add_run(line)
        _set_run_font(run, font=font, size=size, bold=(bold_first and line is lines[0]))
